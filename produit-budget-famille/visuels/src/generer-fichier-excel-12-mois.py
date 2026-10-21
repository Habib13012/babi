# Budget Famille Sans Stress : même charte que les visuels Etsy (crème, vert, corail)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
import datetime as dt

OUT = "/home/user/babi/produit-budget-famille/Budget-Famille-Sans-Stress-12-mois.xlsx"
FONT = "Arial"
TEAL, TEAL2, CREAM, INK, MUTED = "2E7D6B", "E3F2EE", "FBF6EE", "1F2A2E", "5E6B70"
CORAL, SUN, SUN2, CARD, WHITE = "E9785A", "FFF7D6", "FFFBEA", "F7F7F5", "FFFFFF"
RED, REDINK, ORANGE, GREEN = "F8D7DA", "9C0006", "FFE8CC", "D4EDDA"
EUR = '#,##0.00 €;-#,##0.00 €;"-"'
PCT = '0%;-0%;"-"'
line = Side(style="thin", color="ECE6DA")
soft = Border(bottom=line)
box = Border(left=line, right=line, top=line, bottom=line)

def fill(c): return PatternFill("solid", fgColor=c)
def font(size=10, bold=False, color=INK, italic=False): return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)

def page(ws, widths, rows=60):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    for col, w in widths.items(): ws.column_dimensions[col].width = w
    last = max(len(widths) + 2, 22)
    for r in range(1, rows + 1):
        for c in range(1, last + 1): ws.cell(row=r, column=c).fill = fill(CREAM)

def title(ws, text, sub=None, row=2):
    ws.cell(row=row, column=2, value=text).font = font(20, True, TEAL)
    ws.row_dimensions[row].height = 34
    if sub:
        ws.cell(row=row + 1, column=2, value=sub).font = font(11, color=MUTED)
        ws.row_dimensions[row + 1].height = 20

def header(ws, row, col, labels, height=26):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.font = font(10, True, WHITE); c.fill = fill(TEAL)
        c.alignment = Alignment(horizontal="left" if i == 0 else "center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = height

def cellstyle(c, bg=WHITE, fmt=None, f=None, align=None):
    c.fill = fill(bg); c.border = soft; c.font = f or font(10)
    c.alignment = Alignment(vertical="center", horizontal=align)
    if fmt: c.number_format = fmt

def card(ws, row, col1, col2, label, formula, fmt=None, good=False):
    """Carte KPI : libellé (petit, gris) au-dessus, valeur (grande, grasse)."""
    for r in (row, row + 1):
        if col2 > col1: ws.merge_cells(start_row=r, start_column=col1, end_row=r, end_column=col2)
        for c in range(col1, col2 + 1):
            cell = ws.cell(row=r, column=c); cell.fill = fill(TEAL2 if good else CARD)
            cell.border = Border(left=line if c == col1 else None, right=line if c == col2 else None,
                                 top=line if r == row else None, bottom=line if r == row + 1 else None)
    a = ws.cell(row=row, column=col1, value=label); a.font = font(9, True, MUTED)
    a.alignment = Alignment(vertical="bottom", indent=1)
    b = ws.cell(row=row + 1, column=col1, value=formula); b.font = font(15, True, TEAL if good else INK)
    b.alignment = Alignment(vertical="top", horizontal="left", indent=1)
    if fmt: b.number_format = fmt
    ws.row_dimensions[row].height = 20; ws.row_dimensions[row + 1].height = 28

ABR = ["Janv.", "Févr.", "Mars", "Avr.", "Mai", "Juin", "Juil.", "Août", "Sept.", "Oct.", "Nov.", "Déc."]
MOIS = ["Janvier", "Février", "Mars", "Avril", "Mai", "Juin", "Juillet", "Août", "Septembre", "Octobre", "Novembre", "Décembre"]
wb = Workbook()

# ---------------- Mode d'emploi ----------------
m = wb.active; m.title = "Mode d'emploi"
page(m, {"B": 6, "C": 58})
m["B2"] = "ÉDITION 12 MOIS"; m["B2"].font = font(9, True, WHITE)
m.merge_cells("B2:C2"); m["B2"].fill = fill(CORAL); m["B2"].alignment = Alignment(indent=1, vertical="center"); m.row_dimensions[2].height = 22
m["B4"] = "Budget Famille Sans Stress"; m["B4"].font = font(24, True, TEAL); m.merge_cells("B4:C4"); m.row_dimensions[4].height = 40
m["B5"] = "Toute votre année de budget familial dans un seul fichier. Vous notez, il calcule tout."
m["B5"].font = font(11, color=MUTED); m["B5"].alignment = Alignment(wrap_text=True, vertical="center"); m.merge_cells("B5:C5"); m.row_dimensions[5].height = 40
def section(r, text):
    m.cell(row=r, column=2, value=text).font = font(14, True, TEAL); m.row_dimensions[r].height = 30
section(7, "Comment ça marche")
steps = [("Paramètres", "indiquez l'année, vos revenus de chaque mois et le budget mensuel de chaque catégorie."),
         ("Dépenses", "notez toutes les dépenses de l'année dans ce seul journal (date, description, catégorie, montant)."),
         ("Tableau de bord", "choisissez un mois dans la liste en haut : tout se recalcule pour ce mois."),
         ("Bilan annuel", "vos 12 mois d'un coup d'œil, catégorie par catégorie, avec les dépassements en rouge."),
         ("Défi 30 jours", "choisissez « Oui » chaque jour réussi et notez ce que vous avez économisé.")]
for i, (tab, txt) in enumerate(steps):
    r = 8 + i
    n = m.cell(row=r, column=2, value=i + 1); n.font = font(12, True, WHITE); n.fill = fill(CORAL)
    n.alignment = Alignment(horizontal="center", vertical="center")
    t = m.cell(row=r, column=3, value=f"Onglet « {tab} » : {txt}"); t.fill = fill(WHITE)
    t.font = font(10); t.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
    m.row_dimensions[r].height = 40
    for c in (2, 3): m.cell(row=r, column=c).border = Border(bottom=Side(style="thin", color=CREAM), top=Side(style="thin", color=CREAM))
section(14, "Légende")
leg = [(SUN, "Case jaune : à remplir par vous"), (WHITE, "Case blanche : calculée automatiquement, ne pas effacer"),
       (ORANGE, "Orange : 80 % du budget atteint"), (RED, "Rouge : budget dépassé")]
for i, (col, txt) in enumerate(leg):
    r = 15 + i
    m.cell(row=r, column=2).fill = fill(col); m.cell(row=r, column=2).border = box
    t = m.cell(row=r, column=3, value=txt); t.font = font(10); t.alignment = Alignment(vertical="center", indent=1)
    m.row_dimensions[r].height = 22
section(20, "Bon à savoir")
tips = ["Les lignes en gris italique dans « Dépenses » sont des exemples : supprimez-les avant de commencer.",
        "Les dates doivent être de vraies dates (ex. 15/01/2027) : c'est elles qui rangent chaque dépense dans le bon mois.",
        "Nouvelle année : Fichier > Faire une copie (Google Sheets) ou Enregistrer sous (Excel), changez l'année dans « Paramètres » et videz « Dépenses ».",
        "Google Sheets (gratuit) : sheets.google.com > Fichier > Importer > choisissez ce fichier.",
        "Sur téléphone, les couleurs s'affichent le mieux avec le thème clair."]
for i, txt in enumerate(tips):
    r = 21 + i
    m.cell(row=r, column=2, value="•").font = font(12, True, CORAL)
    m.cell(row=r, column=2).alignment = Alignment(horizontal="center", vertical="top")
    t = m.cell(row=r, column=3, value=txt); t.font = font(10); t.alignment = Alignment(wrap_text=True, vertical="top")
    m.row_dimensions[r].height = 30

# ---------------- Paramètres ----------------
p = wb.create_sheet("Paramètres")
page(p, {"B": 34, "C": 17, "D": 13}, rows=70)
title(p, "Paramètres de l'année", "Remplissez les cases jaunes en début d'année (modifiables à tout moment).")
def inp(c, fmt=None): cellstyle(c, SUN, fmt)
for r, lab, val, fmt in ((5, "Année", 2027, "0"), (6, "Objectif d'épargne mensuel", 150, EUR)):
    cellstyle(p.cell(row=r, column=2, value=lab), WHITE, f=font(10, True)); inp(p.cell(row=r, column=3, value=val), fmt)
    p.row_dimensions[r].height = 22
header(p, 8, 2, ["Revenus de chaque mois", "Montant"])
REV_FIRST = 9
for i, mo in enumerate(MOIS):
    r = REV_FIRST + i
    cellstyle(p.cell(row=r, column=2, value=mo), WHITE, f=font(10))
    inp(p.cell(row=r, column=3, value=3450), EUR); p.row_dimensions[r].height = 21
REV_LAST = REV_FIRST + 11
r = REV_LAST + 1
cellstyle(p.cell(row=r, column=2, value="Total revenus de l'année"), TEAL2, f=font(10, True, TEAL))
cellstyle(p.cell(row=r, column=3, value=f"=SUM(C{REV_FIRST}:C{REV_LAST})"), TEAL2, EUR, font(10, True, TEAL)); p.row_dimensions[r].height = 24
p.cell(row=r + 1, column=2, value="Exemple : 1 800 € + 1 500 € de salaires + 150 € de CAF = 3 450 €. Mettez vos vrais montants.").font = font(9, italic=True, color=MUTED)
H2 = r + 3
header(p, H2, 2, ["Catégorie de dépense", "Budget mensuel", "Type"])
cats = [("Loyer / crédit immobilier", 900, "Fixe"), ("Électricité, gaz, eau", 140, "Fixe"),
        ("Internet et téléphones", 50, "Fixe"), ("Assurances", 80, "Fixe"),
        ("Crèche / nounou / garde", 250, "Fixe"), ("Cantine et périscolaire", 80, "Fixe"),
        ("Courses alimentaires", 550, "Variable"), ("Couches, hygiène, puériculture", 60, "Variable"),
        ("Vêtements enfants", 50, "Variable"), ("Activités et loisirs enfants", 40, "Variable"),
        ("Santé / pharmacie", 30, "Variable"), ("Transport / carburant", 150, "Variable"),
        ("Sorties et loisirs adultes", 60, "Variable"), ("Cadeaux et anniversaires", 30, "Variable"),
        ("Abonnements (streaming, etc.)", 25, "Variable"), ("Imprévus", 50, "Variable"),
        (None, None, None), (None, None, None), (None, None, None), (None, None, None)]
CF = H2 + 1; CL = CF + len(cats) - 1
for i, (n, v, ty) in enumerate(cats):
    r = CF + i
    inp(p.cell(row=r, column=2, value=n)); inp(p.cell(row=r, column=3, value=v), EUR)
    inp(p.cell(row=r, column=4, value=ty)); p.cell(row=r, column=4).alignment = Alignment(horizontal="center", vertical="center")
    p.row_dimensions[r].height = 22
dvt = DataValidation(type="list", formula1='"Fixe,Variable"', allow_blank=True); p.add_data_validation(dvt); dvt.add(f"D{CF}:D{CL}")
r = CL + 1
cellstyle(p.cell(row=r, column=2, value="Total budget mensuel"), TEAL2, f=font(10, True, TEAL))
cellstyle(p.cell(row=r, column=3, value=f"=SUM(C{CF}:C{CL})"), TEAL2, EUR, font(10, True, TEAL)); cellstyle(p.cell(row=r, column=4), TEAL2)
p.row_dimensions[r].height = 24
p.cell(row=r + 2, column=2, value="Montants d'exemple : remplacez-les par les vôtres. 4 lignes vides pour vos propres catégories.").font = font(9, italic=True, color=MUTED)
ANNEE = "'Paramètres'!$C$5"
CAT_RANGE = f"'Paramètres'!$B${CF}:$B${CL}"

# ---------------- Dépenses (journal annuel) ----------------
d = wb.create_sheet("Dépenses")
N = 1500
page(d, {"B": 13, "C": 30, "D": 30, "E": 14, "F": 15}, rows=N + 10)
title(d, "Journal des dépenses de l'année", "Une ligne par dépense, toute l'année. Le mois est déduit de la date.")
header(d, 5, 2, ["Date", "Description", "Catégorie", "Montant", "Payé par"])
ex = [(dt.date(2027, 1, 2), "Courses supermarché", "Courses alimentaires", 87.40, "Parent 1"),
      (dt.date(2027, 1, 3), "Paquet de couches", "Couches, hygiène, puériculture", 18.90, "Parent 2"),
      (dt.date(2027, 1, 5), "Loyer janvier", "Loyer / crédit immobilier", 900, "Compte commun"),
      (dt.date(2027, 1, 6), "Facture électricité", "Électricité, gaz, eau", 98.40, "Compte commun"),
      (dt.date(2027, 1, 9), "Cours de natation", "Activités et loisirs enfants", 25, "Parent 1"),
      (dt.date(2027, 2, 2), "Courses supermarché", "Courses alimentaires", 112.60, "Parent 2"),
      (dt.date(2027, 2, 5), "Loyer février", "Loyer / crédit immobilier", 900, "Compte commun")]
for r in range(6, 6 + N):
    for c in range(2, 7): cellstyle(d.cell(row=r, column=c), SUN if r % 2 == 0 else SUN2)
    d.cell(row=r, column=2).number_format = "dd/mm/yyyy"; d.cell(row=r, column=5).number_format = EUR
    d.row_dimensions[r].height = 21
for i, row in enumerate(ex):
    for j, v in enumerate(row):
        c = d.cell(row=6 + i, column=2 + j, value=v); c.font = font(10, color=MUTED, italic=True)
dv = DataValidation(type="list", formula1=f"={CAT_RANGE}", allow_blank=True, showErrorMessage=True,
                    errorTitle="Catégorie inconnue", error="Choisissez une catégorie de la liste (onglet Paramètres).")
d.add_data_validation(dv); dv.add(f"D6:D{5 + N}")
dvd = DataValidation(type="date", operator="greaterThan", formula1="DATE(2000,1,1)", allow_blank=True, showErrorMessage=True,
                     errorTitle="Date invalide", error="Saisissez une date, par exemple 15/01/2027.")
d.add_data_validation(dvd); dvd.add(f"B6:B{5 + N}")
d.freeze_panes = "A6"
DEP_DATE, DEP_CAT, DEP_MT = f"'Dépenses'!$B$6:$B${5 + N}", f"'Dépenses'!$D$6:$D${5 + N}", f"'Dépenses'!$E$6:$E${5 + N}"
def periode(start, end): return f'{DEP_DATE},">="&{start},{DEP_DATE},"<"&{end}'

# ---------------- Tableau de bord (mois au choix) ----------------
t = wb.create_sheet("Tableau de bord")
page(t, {"B": 34, "C": 15, "D": 15, "E": 15, "F": 12, "G": 16}, rows=70)
title(t, "Tableau de bord du mois")
cellstyle(t["B3"], WHITE, f=font(10, True, MUTED)); t["B3"] = "Choisissez le mois ▸"
t["C3"] = "Janvier"; cellstyle(t["C3"], SUN, f=font(11, True, TEAL), align="center"); t.row_dimensions[3].height = 24
dvm = DataValidation(type="list", formula1=f"='Paramètres'!$B${REV_FIRST}:$B${REV_LAST}", allow_blank=False); t.add_data_validation(dvm); dvm.add("C3")
t["D3"] = f'={ANNEE}'; t["D3"].font = font(11, True, MUTED); t["D3"].number_format = "0"; t["D3"].alignment = Alignment(horizontal="left", vertical="center")
# cellules techniques (colonne Z masquée) : n° du mois, début et fin de période
t["Z1"] = f"=IFERROR(MATCH($C$3,'Paramètres'!$B${REV_FIRST}:$B${REV_LAST},0),1)"
t["Z2"] = f"=DATE({ANNEE},$Z$1,1)"; t["Z3"] = f"=DATE({ANNEE},$Z$1+1,1)"
t.column_dimensions["Z"].hidden = True
S, E = "$Z$2", "$Z$3"
card(t, 5, 2, 2, "Revenus du mois", f"=INDEX('Paramètres'!$C${REV_FIRST}:$C${REV_LAST},$Z$1)", EUR)
card(t, 5, 3, 4, "Dépensé ce mois", f"=SUMIFS({DEP_MT},{periode(S, E)})", EUR)
card(t, 5, 5, 7, "Reste à vivre", "=B6-C6", EUR, good=True)
card(t, 8, 2, 2, "Objectif d'épargne", "='Paramètres'!C6", EUR)
card(t, 8, 3, 7, "Objectif atteint ?", '=IF(E6>=B9,"Atteint ✓","Pas encore : courage !")', good=True)
t.row_dimensions[7].height = 8; t.row_dimensions[10].height = 12
H = 11
header(t, H, 2, ["Catégorie", "Budget du mois", "Dépensé", "Reste", "% utilisé", "Progression"])
first = H + 1
for k in range(len(cats)):
    r, pr = first + k, CF + k
    vals = [f"=IF('Paramètres'!B{pr}=\"\",\"\",'Paramètres'!B{pr})",
            f"=IF(B{r}=\"\",\"\",'Paramètres'!C{pr})",
            f"=IF(B{r}=\"\",\"\",SUMIFS({DEP_MT},{DEP_CAT},B{r},{periode(S, E)}))",
            f"=IF(B{r}=\"\",\"\",C{r}-D{r})",
            f"=IF(OR(B{r}=\"\",N(C{r})=0),\"\",D{r}/C{r})",
            f'=IF(F{r}="","",REPT("■",MIN(10,ROUND(F{r}*10,0)))&REPT("□",10-MIN(10,ROUND(F{r}*10,0))))']
    for j, v in enumerate(vals):
        c = t.cell(row=r, column=2 + j, value=v)
        cellstyle(c, WHITE, PCT if j == 4 else (EUR if 0 < j < 4 else None),
                  f=font(10, color=TEAL) if j == 5 else None, align=None if j in (0, 5) else "right")
    t.row_dimensions[r].height = 22
last = first + len(cats) - 1; tot = last + 1
cellstyle(t.cell(row=tot, column=2, value="Total"), TEAL2, f=font(10, True, TEAL))
for j, L in ((3, "C"), (4, "D"), (5, "E")):
    cellstyle(t.cell(row=tot, column=j, value=f"=SUM({L}{first}:{L}{last})"), TEAL2, EUR, font(10, True, TEAL), "right")
cellstyle(t.cell(row=tot, column=6, value=f'=IF(C{tot}=0,"",D{tot}/C{tot})'), TEAL2, PCT, font(10, True, TEAL), "right")
cellstyle(t.cell(row=tot, column=7), TEAL2); t.row_dimensions[tot].height = 24
t.cell(row=tot + 1, column=2, value="Dépenses du mois sans catégorie reconnue").font = font(9, italic=True, color=MUTED)
c = t.cell(row=tot + 1, column=4, value=f"=C6-D{tot}"); c.number_format = EUR; c.font = font(9, italic=True, color=MUTED); c.alignment = Alignment(horizontal="right")
rng = f"B{first}:G{last}"
t.conditional_formatting.add(rng, FormulaRule(formula=[f'AND($B{first}<>"",N($F{first})>1)'], fill=fill(RED), font=Font(name=FONT, color=REDINK, bold=True)))
t.conditional_formatting.add(rng, FormulaRule(formula=[f'AND($B{first}<>"",N($F{first})>=0.8,N($F{first})<=1)'], fill=fill(ORANGE)))
t.cell(row=tot + 3, column=2, value="Orange = 80 % du budget atteint · Rouge = budget dépassé").font = font(9, italic=True, color=MUTED)
t.freeze_panes = "C12"

# ---------------- Bilan annuel ----------------
b = wb.create_sheet("Bilan annuel")
widths = {"B": 30}
for i in range(12): widths[chr(ord("C") + i)] = 11
widths.update({"O": 13, "P": 14, "Q": 13})
page(b, widths, rows=60)
for col in range(18, 23):
    for r in range(1, 61): b.cell(row=r, column=col).fill = fill(CREAM)
title(b, "Bilan annuel", "Vos 12 mois d'un coup d'œil. Rouge = mois où la catégorie a dépassé son budget mensuel. Le reste à vivre s'affiche dès qu'un mois a des dépenses.")
b["Q3"] = f"={ANNEE}"; b["Q3"].number_format = "0"; b["Q3"].font = font(14, True, TEAL); b["Q3"].alignment = Alignment(horizontal="right")
def mcol(i): return chr(ord("C") + i)
def bounds(i): return f"DATE({ANNEE},{i + 1},1)", f"DATE({ANNEE},{i + 2},1)"
header(b, 5, 2, ["Synthèse"] + ABR + ["Total année", "", ""])
for col in (16, 17): b.cell(row=5, column=col).value = None
rows = [("Revenus", lambda i: f"='Paramètres'!C{REV_FIRST + i}"),
        ("Dépenses", lambda i: f"=SUMIFS({DEP_MT},{periode(*bounds(i))})"),
        ("Reste à vivre", lambda i: f"=IF({mcol(i)}7=0,\"\",{mcol(i)}6-{mcol(i)}7)")]
for k, (lab, fx) in enumerate(rows):
    r = 6 + k; good = lab == "Reste à vivre"
    bg = TEAL2 if good else WHITE; fnt = font(10, True, TEAL if good else INK)
    cellstyle(b.cell(row=r, column=2, value=lab), bg, f=fnt)
    for i in range(12): cellstyle(b.cell(row=r, column=3 + i, value=fx(i)), bg, '#,##0 €;-#,##0 €;"-"', fnt, "right")
    cellstyle(b.cell(row=r, column=15, value=f"=SUM(C{r}:N{r})"), bg, '#,##0 €;-#,##0 €;"-"', font(10, True, TEAL if good else INK), "right")
    b.row_dimensions[r].height = 22
b.conditional_formatting.add("C8:O8", FormulaRule(formula=['C8<0'], fill=fill(RED), font=Font(name=FONT, color=REDINK, bold=True)))
H3 = 10
header(b, H3, 2, ["Catégorie"] + ABR + ["Total année", "Budget année", "Écart"])
bf = H3 + 1
for k in range(len(cats)):
    r, pr = bf + k, CF + k
    cellstyle(b.cell(row=r, column=2, value=f"=IF('Paramètres'!B{pr}=\"\",\"\",'Paramètres'!B{pr})"), WHITE)
    for i in range(12):
        s0, s1 = bounds(i)
        cellstyle(b.cell(row=r, column=3 + i, value=f"=IF($B{r}=\"\",\"\",SUMIFS({DEP_MT},{DEP_CAT},$B{r},{periode(s0, s1)}))"),
                  WHITE, '#,##0 €;-#,##0 €;"-"', align="right")
    cellstyle(b.cell(row=r, column=15, value=f"=IF($B{r}=\"\",\"\",SUM(C{r}:N{r}))"), WHITE, '#,##0 €;-#,##0 €;"-"', font(10, True), "right")
    cellstyle(b.cell(row=r, column=16, value=f"=IF($B{r}=\"\",\"\",N('Paramètres'!C{pr})*12)"), WHITE, '#,##0 €;-#,##0 €;"-"', align="right")
    cellstyle(b.cell(row=r, column=17, value=f"=IF($B{r}=\"\",\"\",P{r}-O{r})"), WHITE, '#,##0 €;-#,##0 €;"-"', align="right")
    b.row_dimensions[r].height = 21
bl = bf + len(cats) - 1; bt = bl + 1
cellstyle(b.cell(row=bt, column=2, value="Total"), TEAL2, f=font(10, True, TEAL))
for col in range(3, 18):
    L = chr(ord("A") + col - 1)
    cellstyle(b.cell(row=bt, column=col, value=f"=SUM({L}{bf}:{L}{bl})"), TEAL2, '#,##0 €;-#,##0 €;"-"', font(10, True, TEAL), "right")
b.row_dimensions[bt].height = 24
b.conditional_formatting.add(f"C{bf}:N{bl}", FormulaRule(formula=[f'AND($B{bf}<>"",N(C{bf})>N($P{bf})/12)'], fill=fill(RED), font=Font(name=FONT, color=REDINK, bold=True)))
b.conditional_formatting.add(f"Q{bf}:Q{bt}", FormulaRule(formula=[f'N(Q{bf})<0'], fill=fill(RED), font=Font(name=FONT, color=REDINK, bold=True)))
b.cell(row=bt + 2, column=2, value="Écart = budget de l'année − dépensé. Négatif (rouge) = budget annuel dépassé.").font = font(9, italic=True, color=MUTED)
b.freeze_panes = "C6"

# ---------------- Défi 30 jours ----------------
q = wb.create_sheet("Défi 30 jours")
page(q, {"B": 7, "C": 50, "D": 14, "E": 15})
title(q, "Défi 30 jours d'économies en famille", "Un petit geste par jour : choisissez « Oui » et notez l'économie estimée.")
card(q, 5, 2, 3, "Jours réussis", '=COUNTIF(D11:D40,"Oui")&" / 30"', good=True)
card(q, 5, 4, 5, "Total économisé", "=SUM(E11:E40)", EUR, good=True)
q.row_dimensions[7].height = 8
q.merge_cells("B8:E8"); q["B8"] = "BONUS OFFERT"; q["B8"].font = font(9, True, WHITE); q["B8"].fill = fill(CORAL)
q["B8"].alignment = Alignment(indent=1, vertical="center"); q.row_dimensions[8].height = 20; q.row_dimensions[9].height = 8
defis = ["Faire l'inventaire du frigo et des placards avant les courses", "Préparer une liste de courses et s'y tenir",
    "Repas « fond de placard » ce soir", "Résilier ou suspendre un abonnement peu utilisé", "Journée sans aucune dépense",
    "Préparer le goûter des enfants maison", "Comparer son forfait mobile et internet", "Vendre un vêtement ou jouet trop petit (Vinted, Leboncoin)",
    "Batch cooking : cuisiner 3 repas d'avance", "Sortie gratuite en famille (parc, bibliothèque, balade)", "Vérifier ses droits CAF sur caf.fr (simulateur)",
    "Acheter en marque distributeur pour 5 produits", "Éteindre les appareils en veille de la maison", "Journée sans aucune dépense",
    "Emprunter au lieu d'acheter (livre, jeu, outil)", "Pique-nique au lieu du restaurant", "Utiliser une appli anti-gaspi (ex. Too Good To Go)",
    "Faire le point sur ses assurances", "Douches plus courtes pour toute la famille", "Journée sans aucune dépense",
    "Organiser un troc de vêtements enfants entre parents", "Cuisiner avec les restes", "Désinstaller une appli d'achat compulsif",
    "Préparer les repas de la semaine à l'avance", "Comparer le prix au kilo, pas au paquet", "Après-midi jeux de société à la maison",
    "Journée sans aucune dépense", "Revoir le budget courses du mois prochain", "Mettre l'argent économisé de côté (virement épargne)",
    "Faire le bilan du défi en famille et fêter ça (gratuitement !)"]
header(q, 10, 2, ["Jour", "Défi du jour", "Réussi ?", "Économie"])
for i, txt in enumerate(defis):
    r = 11 + i
    cellstyle(q.cell(row=r, column=2, value=i + 1), WHITE, f=font(10, True, TEAL), align="center")
    c = q.cell(row=r, column=3, value=txt); cellstyle(c); c.alignment = Alignment(vertical="center", wrap_text=True)
    cellstyle(q.cell(row=r, column=4), SUN, align="center"); cellstyle(q.cell(row=r, column=5), SUN, EUR, align="right")
    q.row_dimensions[r].height = 24
dvo = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=True); q.add_data_validation(dvo); dvo.add("D11:D40")
q.conditional_formatting.add("B11:E40", FormulaRule(formula=['$D11="Oui"'], fill=fill(GREEN)))
q.freeze_panes = "A11"

AREAS = {"Mode d'emploi": "D30", "Paramètres": "E50", "Dépenses": f"G{5 + N}", "Tableau de bord": "H40", "Bilan annuel": "R35", "Défi 30 jours": "F42"}
for ws in wb.worksheets:
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.page_setup.orientation = "landscape" if ws.title in ("Tableau de bord", "Bilan annuel") else "portrait"
    ws.print_area = f"A1:{AREAS[ws.title]}"
wb.save(OUT); print(OUT)
